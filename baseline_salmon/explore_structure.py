#!/usr/bin/env python
# coding: utf-8
"""
Prints out the actual folder/file layout of the downloaded Salmon Re-ID
dataset, so we build the real eval/data-loading script against the real
structure instead of guessing from the paper/README (which didn't give an
exact schema). Run this once after download_salmon_reid.py finishes, and
paste the output back.

USAGE
-----
python explore_structure.py --data-root data\\salmon_reid
"""

import argparse
from pathlib import Path


def describe_dir(path: Path, max_depth=3, max_items=15, prefix=""):
    if max_depth < 0:
        return
    try:
        entries = sorted(path.iterdir())
    except Exception as e:
        print(f"{prefix}[error reading {path}: {e}]")
        return
    dirs = [e for e in entries if e.is_dir()]
    files = [e for e in entries if e.is_file()]
    print(f"{prefix}{path.name}/  ({len(dirs)} subdirs, {len(files)} files)")
    for f in files[:max_items]:
        size_kb = f.stat().st_size / 1024
        print(f"{prefix}  - {f.name}  ({size_kb:.1f} KB)")
    if len(files) > max_items:
        print(f"{prefix}  ... and {len(files) - max_items} more files")
    for d in dirs[:max_items]:
        describe_dir(d, max_depth=max_depth - 1, max_items=max_items, prefix=prefix + "  ")
    if len(dirs) > max_items:
        print(f"{prefix}  ... and {len(dirs) - max_items} more subdirs")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-root", required=True)
    args = parser.parse_args()
    base = Path(args.data_root)

    print(f"=== Top level of {base} ===")
    describe_dir(base, max_depth=0)

    reid_dir = base / "reid_dataset"
    if reid_dir.exists():
        print(f"\n=== {reid_dir} (depth 2, to see analysis folders + one level in) ===")
        describe_dir(reid_dir, max_depth=2, max_items=20)
    else:
        print(f"\n(no reid_dataset/ folder found under {base} yet -- did extraction finish?)")

    idmatch = base / "a15_a16_idmatch_with_traj_IDs.xlsx"
    if idmatch.exists():
        print(f"\nFound idmatch file: {idmatch} ({idmatch.stat().st_size/1024:.1f} KB)")
        try:
            import openpyxl
            wb = openpyxl.load_workbook(idmatch, read_only=True)
            for ws_name in wb.sheetnames:
                ws = wb[ws_name]
                print(f"  Sheet '{ws_name}': {ws.max_row} rows x {ws.max_column} cols")
                rows = ws.iter_rows(min_row=1, max_row=5, values_only=True)
                for r in rows:
                    print(f"    {r}")
        except ImportError:
            print("  (install openpyxl to preview its contents: pip install openpyxl)")
    else:
        print(f"\n(idmatch xlsx not found at {idmatch})")


if __name__ == "__main__":
    main()
